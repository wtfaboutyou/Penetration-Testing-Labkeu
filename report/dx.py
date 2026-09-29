# -*- coding: utf-8 -*-
"""Helper pembuatan elemen .docx."""

from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xA1, 0x11, 0x11)
GREY = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x14, 0x60, 0x2E)
AMBER = RGBColor(0x8A, 0x6D, 0x00)
ORANGE = RGBColor(0xA1, 0x5C, 0x00)

SEV_COLOR = {
    "CRITICAL": RED, "HIGH": ORANGE, "MEDIUM": AMBER,
    "LOW": GREEN, "None": GREY,
}


def shade(el, hexcolor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    el.append(sh)


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.02
    for lvl, sz in ((1, 15), (2, 12.5), (3, 11), (4, 10.5)):
        s = doc.styles["Heading %d" % lvl]
        s.font.name = "Calibri"
        s.font.size = Pt(sz)
        s.font.color.rgb = ACCENT
        s.font.bold = True
        s.paragraph_format.space_before = Pt(9 if lvl == 1 else 6)
        s.paragraph_format.space_after = Pt(2)
        s.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
    tg = doc.styles["Table Grid"]
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", 20), ("left", 55), ("bottom", 20), ("right", 55)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tpr = tg.element.find(qn("w:tblPr"))
    if tpr is None:
        tpr = OxmlElement("w:tblPr")
        tg.element.append(tpr)
    tpr.append(mar)


def compact_tables(doc):
    """Rapatkan paragraf di dalam sel tabel: hemat tinggi baris tanpa
    menghapus isi."""
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing = 1.0


def H(doc, text, level=1, pagebreak=False):
    if pagebreak:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    return doc.add_heading(text, level=level)


def P(doc, text="", bold=False, italic=False, color=None, size=None,
      align=None, space_after=None, style=None):
    par = doc.add_paragraph(style=style)
    if text:
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
        if size is not None:
            r.font.size = Pt(size)
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def R(doc, parts, style=None, align=None, space_after=None):
    """parts = [(teks, {b,i,c,mono,sz}), ...]"""
    par = doc.add_paragraph(style=style)
    for text, a in parts:
        r = par.add_run(text)
        r.bold = a.get("b", False)
        r.italic = a.get("i", False)
        if "c" in a:
            r.font.color.rgb = a["c"]
        if a.get("mono"):
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        if "sz" in a:
            r.font.size = Pt(a["sz"])
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    return par


def BUL(doc, text, level=0):
    par = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    par.add_run(text)
    par.paragraph_format.space_after = Pt(2)
    return par


def NUM(doc, text):
    par = doc.add_paragraph(style="List Number")
    par.add_run(text)
    par.paragraph_format.space_after = Pt(2)
    return par


def CODE(doc, lines, caption=None):
    if caption:
        c = doc.add_paragraph()
        r = c.add_run(caption)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY
        c.paragraph_format.space_after = Pt(1)
        c.paragraph_format.keep_with_next = True
    if isinstance(lines, str):
        lines = lines.split("\n")
    for i, ln in enumerate(lines):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        pf.keep_with_next = (i < len(lines) - 1)
        r = par.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(7.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        shade(par._p.get_or_add_pPr(), "F2F3F5")
    tiny_par(doc, 4)


def CALLOUT(doc, title, text, color=RED, fill="FDECEA"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c._tc.get_or_add_tcPr(), fill)
    c.text = ""
    par = c.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    r = par.add_run(title + " ")
    r.bold = True
    r.font.color.rgb = color
    r.font.size = Pt(10)
    r2 = par.add_run(text)
    r2.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def TABLE(doc, headers, rows, widths=None, font=9, header_fill="1F4E79", zebra=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell._tc.get_or_add_tcPr(), header_fill)
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.space_before = Pt(1)
        r = par.add_run(str(htxt))
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell = cells[i]
            if zebra and ri % 2 == 1:
                shade(cell._tc.get_or_add_tcPr(), "F4F6F8")
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.space_before = Pt(1)
            txt = str(val)
            r = par.add_run(txt)
            r.font.size = Pt(font)
            if txt in SEV_COLOR:
                r.bold = True
                r.font.color.rgb = SEV_COLOR[txt]
    if widths:
        t.autofit = False
        lay = OxmlElement("w:tblLayout")
        lay.set(qn("w:type"), "fixed")
        t._tbl.tblPr.append(lay)
        grid = t._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(int(Cm(w).twips)))
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    tiny_par(doc, 3)
    return t


def KV(doc, pairs, w=(4.6, 11.4), font=8.5):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in pairs:
        cells = t.add_row().cells
        shade(cells[0]._tc.get_or_add_tcPr(), "EDF1F5")
        cells[0].text = ""
        r0 = cells[0].paragraphs[0].add_run(k)
        r0.bold = True
        r0.font.size = Pt(font)
        cells[1].text = ""
        r1 = cells[1].paragraphs[0].add_run(str(v))
        r1.font.size = Pt(font)
        if str(v) in SEV_COLOR:
            r1.bold = True
            r1.font.color.rgb = SEV_COLOR[str(v)]
        cells[0].width = Cm(w[0])
        cells[1].width = Cm(w[1])
    tiny_par(doc, 3)
    return t


def FOOTER(section, left_text):
    par = section.footer.paragraphs[0]
    par.text = ""
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(left_text + "   |   Halaman ")
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    f1 = OxmlElement("w:fldSimple")
    f1.set(qn("w:instr"), "PAGE")
    par._p.append(f1)
    r2 = par.add_run(" dari ")
    r2.font.size = Pt(8)
    r2.font.color.rgb = GREY
    f2 = OxmlElement("w:fldSimple")
    f2.set(qn("w:instr"), "NUMPAGES")
    par._p.append(f2)


def tiny_par(doc, space_after=0):
    """Paragraf kosong setipis mungkin (2pt).

    Dipakai sebagai pemisah visual antar blok agar tidak memakan halaman.
    """
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "4")
    rPr.append(sz)
    pPr.append(rPr)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    return p


def HEADER(section, text):
    par = section.header.paragraphs[0]
    par.text = ""
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = par.add_run(text)
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    r.italic = True


def TOC(doc):
    par = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-2" \h \z \u')
    inner = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Daftar Isi - buka di Microsoft Word lalu tekan Ctrl+A lalu F9 untuk populate."
    inner.append(t)
    fld.append(inner)
    par._p.append(fld)
